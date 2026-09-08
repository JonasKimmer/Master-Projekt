"""
Generates poster_1_hsrm.pptx and poster_2_hsrm.pptx from the HSRM A0 template.
Poster 1 → Slide 2 of template (2 small left + large right column)
Poster 2 → Slide 6 of template (1 large top + 3 small middle)
"""

import copy
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from lxml import etree
from pptx.oxml.ns import qn

# ── Colors ───────────────────────────────────────────────────────────────────
P1_PRIMARY = RGBColor(0x4F, 0x6E, 0xF7)   # blue
P1_DARK    = RGBColor(0x1A, 0x1D, 0x2E)
P1_WARM    = RGBColor(0xB4, 0x53, 0x09)

P2_PRIMARY = RGBColor(0x0D, 0x94, 0x88)   # teal
P2_DARK    = RGBColor(0x13, 0x4E, 0x4A)
P2_WARN    = RGBColor(0xD9, 0x77, 0x06)

WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT  = RGBColor(0x20, 0x2C, 0x31)


# ── Helpers ───────────────────────────────────────────────────────────────────
def remove_image_placeholders(slide):
    to_rm = [s for s in slide.shapes if 'Bild' in s.name]
    for s in to_rm:
        s._element.getparent().remove(s._element)


def add_image(slide, path, l, t, w, h):
    return slide.shapes.add_picture(path, Cm(l), Cm(t), Cm(w), Cm(h))


def add_box(slide, l, t, w, h, fill, title, body_lines,
            t_sz=26, b_sz=20, t_bold=True):
    """Rounded rectangle with title line + body lines, all white text."""
    shape = slide.shapes.add_shape(1, Cm(l), Cm(t), Cm(w), Cm(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()

    # Round corners via XML
    sp_pr = shape._element.find(qn('p:spPr'))
    old = sp_pr.find(qn('a:prstGeom'))
    if old is not None:
        sp_pr.remove(old)
    sp_pr.insert(
        0,
        etree.fromstring(
            '<a:prstGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
            ' prst="roundRect"><a:avLst>'
            '<a:gd name="adj" fmla="val 30000"/>'
            '</a:avLst></a:prstGeom>'
        ),
    )

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.55)
    tf.margin_right = Cm(0.55)
    tf.margin_top = Cm(0.45)
    tf.margin_bottom = Cm(0.35)

    # Title
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = title
    r0.font.size = Pt(t_sz)
    r0.font.bold = t_bold
    r0.font.color.rgb = WHITE

    # Body
    for line in body_lines:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(b_sz)
        r.font.bold = False
        r.font.color.rgb = WHITE

    return shape


def set_ph_text(slide, name, lines, sz=None, bold=None, color=None):
    """Replace text in a named placeholder; lines = list of strings (one per para)."""
    for shape in slide.shapes:
        if shape.name != name or not shape.has_text_frame:
            continue
        tf = shape.text_frame
        txBody = tf._txBody

        # Remove all existing <a:p> children
        for p in txBody.findall(qn('a:p')):
            txBody.remove(p)

        for i, line in enumerate(lines):
            p_el = etree.SubElement(txBody, qn('a:p'))
            r_el = etree.SubElement(p_el, qn('a:r'))
            rPr = etree.SubElement(r_el, qn('a:rPr'),
                                   {'lang': 'de-DE', 'dirty': '0'})
            if sz:
                rPr.set('sz', str(int(sz * 100)))
            if bold is not None:
                rPr.set('b', '1' if bold else '0')
            if color:
                sf = etree.SubElement(rPr, qn('a:solidFill'))
                etree.SubElement(sf, qn('a:srgbClr'),
                                 {'val': f'{color.rgb:06X}'})
            t_el = etree.SubElement(r_el, qn('a:t'))
            t_el.text = line
        break


# ── POSTER 1 ──────────────────────────────────────────────────────────────────
def build_poster_1():
    prs = Presentation("_Informatik_Poster-Vorlage_DIN-A0_V1.pptx")
    slide = prs.slides[1]   # Slide 2

    remove_image_placeholders(slide)

    # Scatter plots — left column, stacked
    add_image(slide, 'paper1_scatter_frustration.png', 3, 3,  36, 29)
    add_image(slide, 'paper1_scatter_mentale.png',     3, 34, 36, 29)

    # 3 finding boxes — right column
    add_box(
        slide, 41, 3, 42, 18.5, P1_PRIMARY,
        "① Mentale Last – themenunabhängig",
        [
            "Gaming: 62 · Gesundheit: 63 · Stadtplanung: 63 Pkt.",
            "Unterschied: nur 1,7 Punkte.",
            "Fachwissen schützt nicht vor kognitiver Last.",
        ],
        t_sz=24, b_sz=19,
    )
    add_box(
        slide, 41, 23, 42, 18.5, P1_DARK,
        "② Frustration – stark domänenabhängig",
        [
            "Gaming: 47 · Gesundheit: 44 · Stadtplanung: 37 Pkt.",
            "10 Punkte Unterschied – 6× mehr als mentale Last.",
            "Vertrautheit schützt nicht vor Frustration.",
        ],
        t_sz=24, b_sz=19,
    )
    add_box(
        slide, 41, 43, 42, 20, P1_WARM,
        "③ Bearbeitungszeit ≠ Kognitive Last",
        [
            "Korrelation Dauer × Mentale Last:  r = 0,03  (p = 0,82)",
            "Korrelation Dauer × Frustration:  r = 0,56  (p < 0,001)",
            "Zeit sagt Frustration vorher – nicht mentale Anforderung.",
        ],
        t_sz=24, b_sz=19,
    )

    # Text placeholders
    # Modultitel chip
    set_ph_text(slide, 'Textplatzhalter 36',
                ['Masterprojekt · Wirtschaftsinformatik M.Sc.'])

    # Headline (2 paragraphs = 2 lines)
    set_ph_text(slide, 'Textplatzhalter 35',
                ['Kognitive Beanspruchung', 'bei Datenvisualisierungen'])

    # Description (~500 Zeichen)
    desc = (
        "Wie stark beansprucht eine Datenvisualisierung kognitiv – und hängt das vom Thema ab? "
        "18 Versuchspersonen bearbeiteten je 5 Aufgaben in drei Domänen (Gaming, Gesundheit, Stadtplanung) "
        "und bewerteten ihre Beanspruchung mit dem NASA-TLX. Ergänzt wurden Physiodaten (Shimmer3 GSR/PPG) "
        "und Eye-Tracking (Tobii). Die Ergebnisse zeigen: Mentale Last ist nahezu unabhängig vom Thema "
        "(max. 1,7 Pkt. Unterschied). Frustration dagegen variiert deutlich (10 Pkt.). "
        "Bearbeitungszeit korreliert nicht mit mentaler Last (r = 0,03), wohl aber mit Frustration (r = 0,56)."
    )
    set_ph_text(slide, 'Untertitel 18', [desc])

    # Author line
    set_ph_text(
        slide, 'Textplatzhalter 30',
        ['Jonas Kimmer  //  Hochschule RheinMain  //  Wirtschaftsinformatik M.Sc.  //  2026'],
    )

    # Projekt-Rahmen / Betreuung (right side, below description)
    set_ph_text(slide, 'Textplatzhalter 31', ['Betreuung'])
    set_ph_text(slide, 'Textplatzhalter 34', ['M.Sc. Yasmina Tajja / HSRM Wirtschaftsinformatik'])

    keep_only_slide(prs, 1)
    prs.save('poster_1_hsrm.pptx')
    print("Gespeichert: poster_1_hsrm.pptx")


# ── POSTER 2 ──────────────────────────────────────────────────────────────────
def build_poster_2():
    prs = Presentation("_Informatik_Poster-Vorlage_DIN-A0_V1.pptx")
    slide = prs.slides[5]   # Slide 6

    remove_image_placeholders(slide)

    # Cluster scatter — full width, top
    add_image(slide, 'figures/paper2_abb2_cluster_scatter.png', 3, 3, 78, 43)

    # 3 finding boxes — row below scatter
    add_box(
        slide, 3, 48, 25, 16, P2_PRIMARY,
        "① Zwei Merkmale reichen",
        [
            "Linkanzahl + DOM-Tiefe",
            "genügen zur Klassifikation.",
            "Alle anderen zweitrangig.",
        ],
        t_sz=24, b_sz=19,
    )
    add_box(
        slide, 29.5, 48, 25, 16, P2_DARK,
        "② Kein Label von Hand",
        [
            "Normales ML braucht manuelle",
            "Beispiele. Hier klassifiziert",
            "der Algorithmus vollständig selbst.",
        ],
        t_sz=24, b_sz=19,
    )
    add_box(
        slide, 56, 48, 26, 16, P2_WARN,
        "③ Formulare? Fast egal.",
        [
            "Stärkste Interaktion, aber",
            "erklärt Komplexität am wenigsten.",
            "Strukturmerkmale dominieren.",
        ],
        t_sz=24, b_sz=19,
    )

    # Modultitel chip — Slide 6 uses 'Textplatzhalter 200'
    set_ph_text(slide, 'Textplatzhalter 200',
                ['Masterprojekt · Wirtschaftsinformatik M.Sc.'])

    # Headline
    set_ph_text(slide, 'Textplatzhalter 80',
                ['Wie komplex ist eine', 'Website wirklich?'])

    # Description
    desc = (
        "Python · K-Means Clustering · 3 Test-Websites · 47 Unterseiten\n"
        "6 Merkmale: DOM-Tiefe · Knoten · Links · Text · Medien · Formulare\n"
        "Forschungsfrage: Lässt sich Website-Komplexität automatisch messen "
        "und klassifizieren, ohne manuelle Inspektion?\n"
        "Ergebnis: Drei Seitentypen vollautomatisch erkannt. "
        "Zwei Merkmale reichen zur Klassifikation. "
        "Kein einziges manuell vergebenes Label nötig."
    )
    set_ph_text(slide, 'Untertitel 15', desc.split('\n'))

    # Author line
    set_ph_text(
        slide, 'Textplatzhalter 43',
        ['Jonas Kimmer  //  Hochschule RheinMain  //  Wirtschaftsinformatik M.Sc.  //  2026'],
    )

    # Betreuung
    set_ph_text(slide, 'Textplatzhalter 58', ['Seminarkontext'])
    set_ph_text(slide, 'Textplatzhalter 59',
                ['Masterprojekt · Wirtschaftsinformatik M.Sc. · HSRM Wiesbaden'])
    set_ph_text(slide, 'Textplatzhalter 60', ['Betreuung'])
    set_ph_text(slide, 'Textplatzhalter 61', ['M.Sc. Yasmina Tajja / HSRM Wirtschaftsinformatik'])

    keep_only_slide(prs, 5)
    prs.save('poster_2_hsrm.pptx')
    print("Gespeichert: poster_2_hsrm.pptx")


def remove_slide(prs, index):
    """Remove a slide by index from a presentation (XML manipulation)."""
    slide_part = prs.slides[index].part
    prs_part = prs.part
    rId = None
    for rid, rel in prs_part.rels.items():
        if rel.target_part is slide_part:
            rId = rid
            break
    if rId is None:
        return
    prs_part.drop_rel(rId)
    ns_r = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    for sldId in prs.slides._sldIdLst.findall(qn('p:sldId')):
        if sldId.get(f'{{{ns_r}}}id') == rId:
            prs.slides._sldIdLst.remove(sldId)
            break


def keep_only_slide(prs, keep_index):
    """Delete every slide except keep_index."""
    for i in range(len(prs.slides) - 1, -1, -1):
        if i != keep_index:
            remove_slide(prs, i)


if __name__ == '__main__':
    build_poster_1()
    build_poster_2()
